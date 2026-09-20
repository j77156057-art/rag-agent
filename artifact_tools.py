"""Deterministic document artifact generation for the DocMind agent.

The public entry point accepts a compact JSON payload and creates one DOCX,
PDF, PPTX, or XLSX file below the active project's ``artifacts`` directory.
It deliberately exposes a structured renderer instead of arbitrary code so it
also works in the frozen desktop build where ``python_exec`` is unavailable.
"""
from __future__ import annotations

import json
import os
import re
import tempfile
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from config import CODE_ROOT, STATE_ROOT, get_runtime


SUPPORTED_FORMATS = {"docx", "pdf", "pptx", "xlsx"}
MAX_INPUT_CHARS = 300_000
MAX_ITEMS = 100


class ArtifactError(ValueError):
    pass


def _text(value: Any, limit: int = 20_000) -> str:
    if value is None:
        return ""
    return str(value)[:limit]


def _list(value: Any, *, name: str) -> list:
    if value is None:
        return []
    if not isinstance(value, list):
        raise ArtifactError(f"{name} 必须是数组")
    if len(value) > MAX_ITEMS:
        raise ArtifactError(f"{name} 最多允许 {MAX_ITEMS} 项")
    return value


def _output_dir() -> Path:
    root = get_runtime("code_root") or CODE_ROOT or STATE_ROOT
    path = Path(root).resolve() / "artifacts"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _filename(value: Any, fmt: str) -> str:
    raw = _text(value, 120).strip() or f"document.{fmt}"
    if Path(raw).name != raw:
        raise ArtifactError("filename 只能是文件名，不能包含目录")
    stem = Path(raw).stem if Path(raw).suffix else raw
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", stem).strip(" .")
    if not stem:
        stem = "document"
    return f"{stem}.{fmt}"


def _unique_path(directory: Path, filename: str) -> Path:
    candidate = directory / filename
    if not candidate.exists():
        return candidate
    stem, suffix = candidate.stem, candidate.suffix
    for index in range(2, 10_000):
        candidate = directory / f"{stem}-{index}{suffix}"
        if not candidate.exists():
            return candidate
    raise ArtifactError("无法分配输出文件名")


def _sections(data: dict) -> list[dict]:
    sections = _list(data.get("sections"), name="sections")
    normalized = []
    for section in sections:
        if not isinstance(section, dict):
            raise ArtifactError("sections 的每一项必须是对象")
        normalized.append(section)
    return normalized


def _render_docx(data: dict, path: Path) -> dict:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    title = _text(data.get("title"), 500).strip()
    if title:
        paragraph = document.add_heading(title, level=0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = _text(data.get("subtitle"), 1_000).strip()
    if subtitle:
        paragraph = document.add_paragraph(subtitle)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_content(block: dict) -> None:
        heading = _text(block.get("heading"), 1_000).strip()
        if heading:
            level = max(1, min(3, int(block.get("level", 1))))
            document.add_heading(heading, level=level)
        for paragraph in _list(block.get("paragraphs"), name="paragraphs"):
            document.add_paragraph(_text(paragraph))
        for bullet in _list(block.get("bullets"), name="bullets"):
            document.add_paragraph(_text(bullet), style="List Bullet")
        table = block.get("table")
        if table:
            _add_docx_table(document, table)

    add_content(data)
    for block in _sections(data):
        add_content(block)

    document.save(path)
    reopened = Document(path)
    return {"paragraphs": len(reopened.paragraphs), "tables": len(reopened.tables)}


def _add_docx_table(document, table: Any) -> None:
    if not isinstance(table, dict):
        raise ArtifactError("table 必须是对象")
    headers = _list(table.get("headers"), name="table.headers")
    rows = _list(table.get("rows"), name="table.rows")
    width = len(headers) or max((len(row) for row in rows if isinstance(row, list)), default=0)
    if not width:
        return
    widget = document.add_table(rows=1 if headers else 0, cols=width)
    widget.style = "Table Grid"
    if headers:
        for index, value in enumerate(headers[:width]):
            widget.rows[0].cells[index].text = _text(value)
    for row in rows:
        if not isinstance(row, list):
            raise ArtifactError("table.rows 的每一项必须是数组")
        cells = widget.add_row().cells
        for index, value in enumerate(row[:width]):
            cells[index].text = _text(value)


def _register_pdf_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = (
        ("DocMindCJK", r"C:\Windows\Fonts\msyh.ttc"),
        ("DocMindCJK", r"C:\Windows\Fonts\simhei.ttf"),
        ("DocMindCJK", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    )
    for name, font_path in candidates:
        if os.path.isfile(font_path):
            try:
                pdfmetrics.registerFont(TTFont(name, font_path))
                return name
            except Exception:
                continue
    return "Helvetica"


def _render_pdf(data: dict, path: Path) -> dict:
    from pypdf import PdfReader
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (ListFlowable, ListItem, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)

    font = _register_pdf_font()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("DocMindBody", parent=styles["BodyText"], fontName=font,
                          fontSize=10.5, leading=16, wordWrap="CJK", spaceAfter=5)
    title_style = ParagraphStyle("DocMindTitle", parent=body, fontSize=22,
                                 leading=28, alignment=TA_CENTER, spaceAfter=10)
    subtitle_style = ParagraphStyle("DocMindSubtitle", parent=body, fontSize=11,
                                    textColor=colors.HexColor("#555555"),
                                    alignment=TA_CENTER, spaceAfter=14)
    heading_styles = {
        1: ParagraphStyle("DocMindH1", parent=body, fontSize=16, leading=22,
                          spaceBefore=12, spaceAfter=6),
        2: ParagraphStyle("DocMindH2", parent=body, fontSize=13, leading=19,
                          spaceBefore=9, spaceAfter=4),
        3: ParagraphStyle("DocMindH3", parent=body, fontSize=11, leading=17,
                          spaceBefore=7, spaceAfter=3),
    }
    story = []
    title = _text(data.get("title"), 500).strip()
    if title:
        story.append(Paragraph(escape(title), title_style))
    subtitle = _text(data.get("subtitle"), 1_000).strip()
    if subtitle:
        story.append(Paragraph(escape(subtitle), subtitle_style))

    def add_content(block: dict) -> None:
        heading = _text(block.get("heading"), 1_000).strip()
        if heading:
            level = max(1, min(3, int(block.get("level", 1))))
            story.append(Paragraph(escape(heading), heading_styles[level]))
        for paragraph in _list(block.get("paragraphs"), name="paragraphs"):
            story.append(Paragraph(escape(_text(paragraph)).replace("\n", "<br/>"), body))
        bullets = _list(block.get("bullets"), name="bullets")
        if bullets:
            story.append(ListFlowable(
                [ListItem(Paragraph(escape(_text(item)), body)) for item in bullets],
                bulletType="bullet", leftIndent=16,
            ))
            story.append(Spacer(1, 4))
        table = block.get("table")
        if table:
            story.append(_pdf_table(table, body))
            story.append(Spacer(1, 7))

    add_content(data)
    for block in _sections(data):
        add_content(block)

    document = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=18 * mm,
                                 leftMargin=18 * mm, topMargin=18 * mm,
                                 bottomMargin=18 * mm, title=title)
    document.build(story or [Paragraph(" ", body)])
    reader = PdfReader(str(path))
    return {"pages": len(reader.pages)}


def _pdf_table(table: Any, body_style):
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Table, TableStyle

    if not isinstance(table, dict):
        raise ArtifactError("table 必须是对象")
    headers = _list(table.get("headers"), name="table.headers")
    rows = _list(table.get("rows"), name="table.rows")
    matrix = []
    if headers:
        matrix.append([Paragraph(escape(_text(value)), body_style) for value in headers])
    for row in rows:
        if not isinstance(row, list):
            raise ArtifactError("table.rows 的每一项必须是数组")
        matrix.append([Paragraph(escape(_text(value)), body_style) for value in row])
    if not matrix:
        matrix = [[Paragraph(" ", body_style)]]
    widget = Table(matrix, repeatRows=1 if headers else 0, hAlign="LEFT")
    widget.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8BEC6")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF1F4")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))
    return widget


def _render_pptx(data: dict, path: Path) -> dict:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    title = _text(data.get("title"), 500).strip()
    subtitle = _text(data.get("subtitle"), 1_000).strip()
    if title:
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = title
        slide.placeholders[1].text = subtitle

    slides = _list(data.get("slides"), name="slides")
    if not slides:
        slides = [
            {"title": section.get("heading", ""),
             "bullets": section.get("bullets") or section.get("paragraphs") or []}
            for section in _sections(data)
        ]
    for item in slides:
        if not isinstance(item, dict):
            raise ArtifactError("slides 的每一项必须是对象")
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = _text(item.get("title"), 500)
        frame = slide.placeholders[1].text_frame
        frame.clear()
        bullets = _list(item.get("bullets"), name="slides.bullets")
        if not bullets and item.get("body"):
            bullets = [item.get("body")]
        for index, bullet in enumerate(bullets):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = _text(bullet)
            paragraph.level = 0
            paragraph.font.size = Pt(24)
            paragraph.font.name = "Microsoft YaHei"

    if not presentation.slides:
        presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)
    reopened = Presentation(path)
    return {"slides": len(reopened.slides)}


def _render_xlsx(data: dict, path: Path) -> dict:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = _list(data.get("sheets"), name="sheets")
    if not sheets:
        table = data.get("table") or {}
        sheets = [{"name": data.get("title") or "Data",
                   "headers": table.get("headers", []), "rows": table.get("rows", [])}]
    for item in sheets:
        if not isinstance(item, dict):
            raise ArtifactError("sheets 的每一项必须是对象")
        name = re.sub(r"[\\/*?:\[\]]", "_", _text(item.get("name"), 31).strip()) or "Sheet"
        worksheet = workbook.create_sheet(name[:31])
        headers = _list(item.get("headers"), name="sheets.headers")
        rows = _list(item.get("rows"), name="sheets.rows")
        if headers:
            worksheet.append([_xlsx_value(value, data) for value in headers])
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="334155")
                cell.alignment = Alignment(vertical="center")
            worksheet.freeze_panes = "A2"
        for row in rows:
            if not isinstance(row, list):
                raise ArtifactError("sheets.rows 的每一项必须是数组")
            worksheet.append([_xlsx_value(value, data) for value in row])
        for column in worksheet.columns:
            width = min(60, max(10, max((len(_text(cell.value, 200)) for cell in column), default=0) + 2))
            worksheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(path)
    reopened = load_workbook(path, read_only=True, data_only=False)
    result = {"sheets": reopened.sheetnames}
    reopened.close()
    return result


def _xlsx_value(value: Any, data: dict) -> Any:
    if isinstance(value, (int, float, bool)) or value is None:
        return value
    text = _text(value)
    if not data.get("allow_formulas") and text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


_RENDERERS = {
    "docx": _render_docx,
    "pdf": _render_pdf,
    "pptx": _render_pptx,
    "xlsx": _render_xlsx,
}


def create_artifact(arg: str) -> str:
    """Create and validate one office artifact from a JSON object."""
    temporary = None
    try:
        raw = str(arg or "").strip()
        if not raw:
            raise ArtifactError("需要 JSON 参数")
        if len(raw) > MAX_INPUT_CHARS:
            raise ArtifactError("输入内容过大")
        data = json.loads(raw)
        if not isinstance(data, dict):
            raise ArtifactError("参数必须是 JSON 对象")
        fmt = _text(data.get("format"), 10).strip().lower().lstrip(".")
        if fmt not in SUPPORTED_FORMATS:
            raise ArtifactError("format 必须是 docx、pdf、pptx 或 xlsx")
        output_dir = _output_dir()
        target = _unique_path(output_dir, _filename(data.get("filename"), fmt))
        handle = tempfile.NamedTemporaryFile(prefix=".docmind-", suffix=f".{fmt}",
                                             dir=output_dir, delete=False)
        temporary = Path(handle.name)
        handle.close()
        validation = _RENDERERS[fmt](data, temporary)
        os.replace(temporary, target)
        temporary = None
        return json.dumps({
            "ok": True,
            "format": fmt,
            "path": str(target),
            "filename": target.name,
            "bytes": target.stat().st_size,
            "validation": validation,
        }, ensure_ascii=False)
    except (ArtifactError, json.JSONDecodeError, OSError, ImportError, ValueError) as exc:
        return json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False)
    except Exception as exc:  # Keep tool failures inside the tool protocol.
        return json.dumps({"ok": False, "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False)
    finally:
        if temporary is not None:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                pass
